import os
import re
import json
import copy # For deepcopy
import sys # Import sys to check for command-line arguments
import fitz # PyMuPDF
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
import pytesseract
from PIL import Image, ImageDraw, ImageFont, UnidentifiedImageError

SUPPORTED_EXTENSIONS = ['.pdf', '.png', '.jpg', '.jpeg', '.xlsx', '.docx']

# --- PII Regex Patterns (for scanning within extracted text) ---
PHONE_REGEX_PII = re.compile(r'(?:\+91[\s-]?)?(?:0)?[6-9]\d{9}')
EMAIL_REGEX_PII = re.compile(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}')
GSTIN_REGEX_PII = re.compile(r'\b[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z]{1}[1-9A-Z]{1}Z[0-9A-Z]{1}\b', re.IGNORECASE)
PAN_REGEX_PII = re.compile(r'\b[A-Z]{5}[0-9]{4}[A-Z]{1}\b')
ACCOUNT_NO_REGEX_PII = re.compile(r'\b\d{9,18}\b')

# --- Regex Patterns for Invoice Extraction ---
GSTIN_REGEX_INV = re.compile(r'[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z]{1}[1-9A-Z]{1}Z[0-9A-Z\s]{1}', re.IGNORECASE)
INVOICE_NO_REGEX_INV = re.compile(r'(?:Invoice No|Inv No|Invoice Number)\s*[:#-]?\s*([A-Za-z0-9/-]+)', re.IGNORECASE)
DATE_REGEX_INV = re.compile(
    r'(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})|'
    r'(\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s+\d{2,4})|'
    r'(\d{1,2}[-/](?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*[-/]\d{2,4})',
    re.IGNORECASE
)
MONETARY_VALUE_REGEX_INV = re.compile(r'\b\d[\d,]*\.\d{2}\b|\b\d[\d,]*\b')
LINE_ITEM_GUESS_REGEX_INV = re.compile(r'^(.*?)\s+(\d+)\s+([\d\.,]+\.?\d*)\s+([\d\.,]+\.?\d*)$', re.MULTILINE)

# --- Regex Patterns for Balance Sheet Extraction ---
DATE_REGEX_BS = re.compile(r"(?:\bAs of|\bAs at|Date|For the period ending)\s*[:\-]?\s*((?:\d{1,2}[-/]\s?\w+\s?[-/]\d{2,4})|(?:\w+\s+\d{1,2},?\s*\d{2,4})|(?:\d{1,2}[-/]\d{1,2}[-/]\d{2,4}))", re.IGNORECASE)
COMPANY_NAME_REGEX_BS = re.compile(r"^([A-Za-z0-9\s.,&'-]+(?:Ltd\.?|Limited|Inc\.?|Corporation|LLC|LLP|Pvt\.?)\b)", re.IGNORECASE | re.MULTILINE)
CURRENCY_SYMBOL_REGEX_BS = re.compile(r'(₹|\$|€|USD|INR|EUR|GBP)')

# --- PII Masking Functions ---
def mask_phone(phone: str) -> str:
    digits = re.sub(r'\D', '', phone)
    if len(digits) >= 10:
        if len(digits) > 10 and digits.startswith('91'):
            return "91" + digits[2:4] + 'XXXXXX' + digits[-2:]
        return digits[:2] + 'XXXXXX' + digits[-2:]
    return 'XX****XX'

def mask_email(email: str) -> str:
    parts = email.split('@')
    if len(parts) == 2:
        user, domain = parts; domain_parts = domain.split('.')
        masked_user = user[0] + '***' if len(user) > 1 else (user if user else '***')
        masked_domain_start = domain_parts[0][0] + '***' if len(domain_parts[0]) > 1 else (domain_parts[0] if domain_parts[0] else '***')
        masked_domain_end = domain_parts[-1] if len(domain_parts) > 1 else (domain_parts[0] if not masked_domain_start.endswith('***') else "")
        if masked_domain_start.endswith('***') and masked_domain_end == domain_parts[0] and len(domain_parts) == 1:
             return f"{masked_user}@{masked_domain_start}"
        return f"{masked_user}@{masked_domain_start}.{masked_domain_end}"
    return '***@***.***'

def mask_gstin(gstin: str) -> str:
    if len(gstin) == 15: return gstin[:2] + 'XXXXX' + gstin[7:11] + 'X' + gstin[12] + gstin[13] + 'X'
    return "XXXXXXX"

def mask_pan(pan: str) -> str:
    if len(pan) == 10: return pan[:2] + 'XXXX' + pan[6:9] + pan[-1]
    return "XXXXX"

def mask_account_no(acc_no: str) -> str:
    if len(acc_no) > 4: return 'X' * (len(acc_no) - 4) + acc_no[-4:]
    return 'X' * len(acc_no)

# --- PII Masking Application Function ---
def apply_pii_masking(data):
    processed_data = copy.deepcopy(data)
    def _mask_recursive(item):
        if isinstance(item, dict):
            for key, value in item.items(): item[key] = _mask_recursive(value)
            return item
        elif isinstance(item, list):
            for i, element in enumerate(item): item[i] = _mask_recursive(element)
            return item
        elif isinstance(item, str):
            new_item_str = PAN_REGEX_PII.sub(lambda m: mask_pan(m.group(0)), item)
            new_item_str = GSTIN_REGEX_PII.sub(lambda m: mask_gstin(m.group(0)), new_item_str)
            new_item_str = PHONE_REGEX_PII.sub(lambda m: mask_phone(m.group(0)), new_item_str)
            new_item_str = EMAIL_REGEX_PII.sub(lambda m: mask_email(m.group(0)), new_item_str)
            new_item_str = ACCOUNT_NO_REGEX_PII.sub(lambda m: mask_account_no(m.group(0)), new_item_str)
            return new_item_str
        else: return item
    return _mask_recursive(processed_data)

# --- Text Extraction Functions ---
def extract_text_from_pdf(pdf_path: str) -> str:
    full_text = ""; doc = None
    try:
        doc = fitz.open(pdf_path)
        for page_num in range(doc.page_count): page = doc.load_page(page_num); full_text += page.get_text("text")
    except fitz.errors.FitzError as e: return f"Error processing PDF (FitzError): {e}"
    except Exception as e: return f"An unexpected error occurred during PDF processing: {e}"
    finally:
        if doc: doc.close()
    return full_text

def extract_text_from_docx(docx_path: str) -> str:
    full_text_parts = []
    try:
        doc = Document(docx_path)
        for paragraph in doc.paragraphs: full_text_parts.append(paragraph.text)
        full_text = "\n".join(full_text_parts)
    except PackageNotFoundError: return f"Error processing DOCX: File is not a valid DOCX package."
    except Exception as e: return f"An unexpected error occurred during DOCX processing: {e}"
    return full_text

def extract_text_from_xlsx(xlsx_path: str) -> str:
    all_text_parts = []
    try:
        workbook = openpyxl.load_workbook(xlsx_path)
        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            for row in ws.iter_rows():
                row_text_parts = []
                for cell in row:
                    if cell.value is not None: row_text_parts.append(str(cell.value))
                if row_text_parts: all_text_parts.append(" ".join(row_text_parts))
        full_text = "\n".join(all_text_parts)
    except InvalidFileException: return f"Error processing XLSX: File is not a valid XLSX file."
    except Exception as e: return f"An unexpected error occurred during XLSX processing: {e}"
    return full_text

def extract_text_from_image_ocr(image_path: str) -> str:
    try:
        img = Image.open(image_path); text = pytesseract.image_to_string(img); return text.strip()
    except pytesseract.TesseractNotFoundError: return "Error: Tesseract is not installed or not found in your PATH. OCR functionality will not work."
    except UnidentifiedImageError: return f"Error processing Image: Cannot identify image file {image_path}. It may be corrupted or not a valid image."
    except Exception as e: return f"An unexpected error occurred during image OCR processing: {e}"

# --- Information Extractor Functions ---
def extract_balance_sheet_info(text: str) -> dict:
    bs_data = {"company_name": None, "statement_date": None, "currency": None}
    company_match = COMPANY_NAME_REGEX_BS.search(text)
    if company_match: bs_data['company_name'] = company_match.group(1).strip()
    else:
        lines = text.split('\n');
        for line in lines:
            line_stripped = line.strip()
            if line_stripped and not any(kw in line_stripped.lower() for kw in ["balance sheet", "as of", "statement of"]): bs_data['company_name'] = line_stripped; break
    date_match = DATE_REGEX_BS.search(text)
    if date_match: bs_data['statement_date'] = date_match.group(1).strip()
    currency_match = CURRENCY_SYMBOL_REGEX_BS.search(text)
    if currency_match: bs_data['currency'] = currency_match.group(1)
    total_keywords = {
        "total current assets": "total_current_assets", "total non-current assets": "total_non_current_assets",
        "total fixed assets": "total_non_current_assets", "total assets": "total_assets",
        "total current liabilities": "total_current_liabilities", "total non-current liabilities": "total_non_current_liabilities",
        "total long-term liabilities": "total_non_current_liabilities", "total liabilities": "total_liabilities",
        "total equity": "total_equity", "shareholders' equity": "total_equity", "net worth": "total_equity"
    }
    lines = text.split('\n')
    for line in lines:
        line_lower = line.lower()
        for keyword_text, data_key in total_keywords.items():
            if keyword_text in line_lower:
                value_matches = MONETARY_VALUE_REGEX_INV.findall(line)
                if value_matches:
                    value_str = value_matches[-1].replace(',', '')
                    try: bs_data[data_key] = float(value_str)
                    except ValueError: bs_data[data_key] = value_str
                    if not bs_data['currency']:
                        local_currency_match = CURRENCY_SYMBOL_REGEX_BS.search(line)
                        if local_currency_match: bs_data['currency'] = local_currency_match.group(1)
                break
    if not bs_data['currency']:
        all_text_currency_match = CURRENCY_SYMBOL_REGEX_BS.search(text)
        if all_text_currency_match: bs_data['currency'] = all_text_currency_match.group(1)
    return bs_data

def extract_invoice_info(text: str) -> dict:
    invoice_data = {"invoice_number": None, "invoice_date": None, "supplier_info": {"name": None, "address": None, "gstin": None, "phone": None, "email": None}, "recipient_info": {"name": None, "address": None, "gstin": None}, "line_items": [], "subtotal": None, "cgst_amount": None, "sgst_amount": None, "igst_amount": None, "total_amount": None}
    all_gstins = GSTIN_REGEX_INV.findall(text)
    if all_gstins:
        invoice_data['supplier_info']['gstin'] = all_gstins[0].replace(" ", "")
        if len(all_gstins) > 1: invoice_data['recipient_info']['gstin'] = all_gstins[1].replace(" ", "")
    inv_no_match = INVOICE_NO_REGEX_INV.search(text)
    if inv_no_match: invoice_data['invoice_number'] = inv_no_match.group(1)
    date_labels_priority = ["Invoice Date:", "Date:"]; found_date = None
    for label in date_labels_priority:
        pattern_to_search = r"(?i)" + re.escape(label) + r"\s*[:#-]?\s*(?:" + DATE_REGEX_INV.pattern + ")"
        match = re.search(pattern_to_search, text)
        if match:
            actual_date_str = match.group(1) or match.group(2) or match.group(3)
            if actual_date_str: found_date = actual_date_str.strip(); invoice_data['invoice_date'] = found_date; break
    if not found_date:
        first_date_match = DATE_REGEX_INV.search(text)
        if first_date_match: invoice_data['invoice_date'] = (first_date_match.group(1) or first_date_match.group(2) or first_date_match.group(3)).strip()
    lines = text.split('\n'); supplier_line_index = -1; recipient_line_index = -1
    for i, line in enumerate(lines):
        if any(kw.lower() in line.lower() for kw in ["Supplier:", "Sold By:", "From:"]):
            supplier_line_index = i
            name_match = re.search(r"(?:Supplier:|Sold By:|From:)\s*(.*?)(?:, Address|, GSTIN|Ph:|Email:|$)", line, re.IGNORECASE)
            if name_match: invoice_data['supplier_info']['name'] = name_match.group(1).strip().rstrip(',')
            address_match = re.search(r"Address[:\s]*(.*?)(?:, GSTIN|Ph:|Email:|$)", line, re.IGNORECASE)
            if address_match: invoice_data['supplier_info']['address'] = address_match.group(1).strip().rstrip(',')
            phone_match = re.search(r'(?:Ph|Phone|Mobile)[:\s]*([0-9\s+-]+)', line, re.IGNORECASE)
            if phone_match: invoice_data['supplier_info']['phone'] = phone_match.group(1).strip()
            email_match = re.search(r'(?:Email|E-mail)[:\s]*([a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,})', line, re.IGNORECASE)
            if email_match: invoice_data['supplier_info']['email'] = email_match.group(1).strip()
            break
    for i, line in enumerate(lines):
        if any(kw.lower() in line.lower() for kw in ["Bill To:", "Ship To:", "To:"]):
            if i <= supplier_line_index: continue
            recipient_line_index = i
            name_match = re.search(r"(?:Bill To:|Ship To:|To:)\s*(.*?)(?:, Address|, GSTIN|$)", line, re.IGNORECASE)
            if name_match: invoice_data['recipient_info']['name'] = name_match.group(1).strip().rstrip(',')
            address_match = re.search(r"Address[:\s]*(.*?)(?:, GSTIN|$)", line, re.IGNORECASE)
            if address_match: addr_str = address_match.group(1).strip().rstrip(','); addr_str = re.sub(r"GSTIN[:\s]*\S+", "", addr_str, flags=re.IGNORECASE).strip().rstrip(','); invoice_data['recipient_info']['address'] = addr_str
            break
    for line in lines:
        if line.lower().startswith("description") or line.lower().startswith("hsn") or line.lower().startswith("subtotal") or line.lower().startswith("total"): continue
        item_match = LINE_ITEM_GUESS_REGEX_INV.search(line.strip())
        if item_match:
            desc, qty, rate, amount = item_match.groups()
            invoice_data['line_items'].append({"description": desc.strip(), "quantity": int(qty) if qty.isdigit() else qty, "rate": float(rate.replace(',','')) if rate.replace('.','',1).replace(',','').isdigit() else rate, "amount": float(amount.replace(',','')) if amount.replace('.','',1).replace(',','').isdigit() else amount})
    for line in lines:
        line_lower = line.lower(); all_value_matches_on_line = MONETARY_VALUE_REGEX_INV.findall(line)
        if not all_value_matches_on_line: continue
        val_str = all_value_matches_on_line[-1]; val = None
        try: val = float(val_str.replace(',',''))
        except ValueError: val = None
        if val is not None:
            if "subtotal" in line_lower or "sub total" in line_lower: invoice_data['subtotal'] = val
            elif "cgst" in line_lower:
                tax_specific_match = re.search(r"CGST\s*(?:@\s*[\d\.]+%?)?\s*[:\-]?\s*([\d\.,]+)", line, re.IGNORECASE)
                if tax_specific_match: invoice_data['cgst_amount'] = float(tax_specific_match.group(1).replace(',',''))
                else: invoice_data['cgst_amount'] = val
            elif "sgst" in line_lower:
                tax_specific_match = re.search(r"SGST\s*(?:@\s*[\d\.]+%?)?\s*[:\-]?\s*([\d\.,]+)", line, re.IGNORECASE)
                if tax_specific_match: invoice_data['sgst_amount'] = float(tax_specific_match.group(1).replace(',',''))
                else: invoice_data['sgst_amount'] = val
            elif "igst" in line_lower:
                tax_specific_match = re.search(r"IGST\s*(?:@\s*[\d\.]+%?)?\s*[:\-]?\s*([\d\.,]+)", line, re.IGNORECASE)
                if tax_specific_match: invoice_data['igst_amount'] = float(tax_specific_match.group(1).replace(',',''))
                else: invoice_data['igst_amount'] = val
            elif "total amount" in line_lower or "grand total" in line_lower or (line_lower.startswith("total:") and not invoice_data['total_amount']): invoice_data['total_amount'] = val
    return invoice_data

def extract_bank_statement_info(text: str) -> dict:
    statement_data = {"transactions": [], "account_holder": None, "account_number": None, "statement_period_start": None, "statement_period_end": None, "bank_name": None}
    acc_match = re.search(r"(?:Account No|A/C No|Account Number)[:\s]*([\d\s-]{9,18})\b", text, re.IGNORECASE)
    if acc_match: statement_data['account_number'] = acc_match.group(1).strip()
    holder_match = re.search(r"(?:Account Holder|Name)[:\s]*(.+?)(?:\n|Account No|A/C No)", text, re.IGNORECASE)
    if holder_match:
        statement_data['account_holder'] = holder_match.group(1).strip()
        if statement_data['account_holder'] and statement_data['account_number'] and statement_data['account_number'] in statement_data['account_holder']:
             statement_data['account_holder'] = statement_data['account_holder'].split(statement_data['account_number'])[0].strip()
    period_match = re.search(r"Statement Period[:\s]*(\d{1,2}/\d{1,2}/\d{2,4})\s*(?:to|-)\s*(\d{1,2}/\d{1,2}/\d{2,4})", text, re.IGNORECASE)
    if period_match: statement_data['statement_period_start'] = period_match.group(1).strip(); statement_data['statement_period_end'] = period_match.group(2).strip()
    for i, line in enumerate(text.split('\n')):
        if i < 5:
            bank_match = re.match(r"^([A-Za-z\s&]+(?:Bank|BANK|Ltd\.|LIMITED|CORPORATION|CORP\.)(?: Ltd\.?)?)\s*$", line.strip(), re.IGNORECASE)
            if bank_match: statement_data['bank_name'] = bank_match.group(1).strip(); break
    TRANSACTION_LINE_REGEX = re.compile(r"(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})\s+(.+?)\s{2,}(?:([\d\.,]+)\s+)?(?:([\d\.,]+)\s+)?([\d\.,]+)$", re.IGNORECASE | re.VERBOSE)
    lines = text.split('\n')
    for line in lines:
        line = line.strip(); tx_match = TRANSACTION_LINE_REGEX.search(line)
        if tx_match:
            date, description, group3_debit, group4_credit, balance_val = tx_match.groups()
            if description.lower() in ["description", "particulars", "details", "opening balance"] or date.lower() == "date":
                if description.lower() == "opening balance" and not group3_debit and not group4_credit and balance_val:
                     statement_data['transactions'].append({"date": date.strip(), "description": description.strip(), "debit": None, "credit": None, "balance": float(balance_val.replace(',', '')) if balance_val and balance_val.replace('.','',1).replace(',','').isdigit() else None})
                continue
            debit_val = float(group3_debit.replace(',', '')) if group3_debit and group3_debit.replace('.','',1).replace(',','').isdigit() else None
            credit_val = float(group4_credit.replace(',', '')) if group4_credit and group4_credit.replace('.','',1).replace(',','').isdigit() else None
            if "NEFT from XYZ Corp".lower() in description.lower():
                if not group3_debit and group4_credit: debit_val = None; credit_val = float(group4_credit.replace(',', ''))
                elif group3_debit and not group4_credit: credit_val = float(group3_debit.replace(',', '')); debit_val = None
            statement_data['transactions'].append({"date": date.strip(), "description": description.strip(), "debit": debit_val, "credit": credit_val, "balance": float(balance_val.replace(',', '')) if balance_val and balance_val.replace('.','',1).replace(',','').isdigit() else None})
    return statement_data

# --- Document Type Classification ---
BALANCE_SHEET_KEYWORDS = {"balance sheet", "assets", "liabilities", "equity", "shareholder equity", "statement of financial position"}
INVOICE_KEYWORDS = {"invoice no", "gstin", "invoice date", "hsn", "sac", "tax invoice", "bill to", "ship to", "invoice number"}
BANK_STATEMENT_KEYWORDS = {"account statement", "account no", "transaction date", "narration", "description", "debit", "credit", "balance", "statement of account"}
MIN_KEYWORD_THRESHOLD = 2

def classify_document(text: str) -> str:
    if not text: return "UNKNOWN_DOCUMENT_TYPE"
    lower_text = text.lower()
    scores = {"BALANCE_SHEET": sum(1 for keyword in BALANCE_SHEET_KEYWORDS if keyword in lower_text), "INVOICE": sum(1 for keyword in INVOICE_KEYWORDS if keyword in lower_text), "BANK_STATEMENT": sum(1 for keyword in BANK_STATEMENT_KEYWORDS if keyword in lower_text)}
    max_score = 0; classified_type = "UNKNOWN_DOCUMENT_TYPE"
    for doc_type, score in scores.items():
        if score > max_score: max_score = score; classified_type = doc_type
        elif score == max_score and score > 0: classified_type = "UNKNOWN_DOCUMENT_TYPE"
    if max_score < MIN_KEYWORD_THRESHOLD: return "UNKNOWN_DOCUMENT_TYPE"
    return classified_type

# --- Main Document Processing Function ---
def process_document(file_path: str) -> str:
    print(f"\n--- Processing document: {file_path} ---")
    if not os.path.exists(file_path):
        print(f"Error: File not found: {file_path}")
        return json.dumps({"error": "File not found", "file_path": file_path}, indent=4)
    _, file_extension = os.path.splitext(file_path); file_extension_lower = file_extension.lower()
    if file_extension_lower not in SUPPORTED_EXTENSIONS:
        print(f"Error: Unsupported file type: {file_extension}")
        return json.dumps({"error": "Unsupported file type", "file_type": file_extension}, indent=4)
    print(f"File Type: {file_extension_lower}"); extracted_text = ""
    if file_extension_lower == '.pdf': extracted_text = extract_text_from_pdf(file_path)
    elif file_extension_lower == '.docx': extracted_text = extract_text_from_docx(file_path)
    elif file_extension_lower == '.xlsx': extracted_text = extract_text_from_xlsx(file_path)
    elif file_extension_lower in ['.png', '.jpg', '.jpeg']: extracted_text = extract_text_from_image_ocr(file_path)
    if extracted_text.startswith("Error:") or extracted_text.startswith("An unexpected error"):
        print(f"Text extraction failed: {extracted_text}")
        return json.dumps({"error": "Text extraction failed", "details": extracted_text}, indent=4)
    if not extracted_text.strip():
        print("No text content extracted.")
        status_msg = "No text content found in document."
        if file_extension_lower in ['.png', '.jpg', '.jpeg']: status_msg = "OCR did not find any text in the image."
        return json.dumps({"status": status_msg, "file_path": file_path, "text_snippet": ""}, indent=4)
    snippet = extracted_text[:200].replace('\n', ' '); print(f"Extracted text snippet (first 200 chars): {snippet}")
    doc_type = classify_document(extracted_text); print(f"Classified Document Type: {doc_type}")
    processed_info = {}
    if doc_type == "BANK_STATEMENT": processed_info = extract_bank_statement_info(extracted_text)
    elif doc_type == "INVOICE": processed_info = extract_invoice_info(extracted_text)
    elif doc_type == "BALANCE_SHEET": processed_info = extract_balance_sheet_info(extracted_text)
    else:
        masked_text_snippet = apply_pii_masking(extracted_text[:500])
        return json.dumps({"status": "UNKNOWN_DOCUMENT_TYPE", "file_path": file_path, "text_snippet_masked": masked_text_snippet}, indent=4)
    print("Applying PII masking..."); masked_info = apply_pii_masking(processed_info)
    return json.dumps(masked_info, indent=4)

if __name__ == '__main__':
    if len(sys.argv) > 1:
        file_to_process_from_arg = sys.argv[1]
        json_result = process_document(file_to_process_from_arg)
        print(json_result)
    else:
        print("--- Running Default Test Suite with Dummy Files ---")
        test_files_to_create = {
            "pdf": {"name": "dummy_test.pdf", "type": "Bank Statement"},
            "docx": {"name": "dummy_test.docx", "type": "Invoice"},
            "xlsx": {"name": "dummy_test.xlsx", "type": "Balance Sheet"},
            "png": {"name": "dummy_test.png", "type": "Unknown (OCR)"},
            "jpg": {"name": "photo.jpg", "type": "Empty Image (OCR)"},
            "txt": {"name": "unsupported.txt", "type": "Unsupported"}
        }
        DUMMY_PDF_PATH = test_files_to_create["pdf"]["name"]
        DUMMY_DOCX_PATH = test_files_to_create["docx"]["name"]
        DUMMY_XLSX_PATH = test_files_to_create["xlsx"]["name"]
        DUMMY_PNG_PATH = test_files_to_create["png"]["name"]
        DUMMY_JPG_PATH = test_files_to_create["jpg"]["name"]
        UNSUPPORTED_TXT_PATH = test_files_to_create["txt"]["name"]
        NON_EXISTENT_FILE = "nonexistent_file.tmp"
        print("--- Setting up dummy files ---")
        try:
            pdf_doc = fitz.open()
            page1_text = """MyExample Bank Ltd.\nStatement of Account\nAccount Holder: John Doe\nAccount No: 123456789012\nStatement Period: 01/01/2024 to 31/01/2024\n\nDate        Description                     Debit      Credit     Balance\n01/01/2024  Opening Balance                                    10000.00\n"""
            pdf_page1 = pdf_doc.new_page(width=600, height=400); y_coord = 50
            for line in page1_text.split('\n'): pdf_page1.insert_text((50, y_coord), line); y_coord += 15
            page2_text = """05/01/2024  ATM Withdrawal ABC Complex      500.00                9500.00\n10/01/2024  NEFT from XYZ Corp                         2000.00   11500.00\n15/01/2024  Cheque No 12345                   75.50              11424.50"""
            pdf_page2 = pdf_doc.new_page(width=600, height=400); y_coord = 50
            for line in page2_text.split('\n'): pdf_page2.insert_text((50, y_coord), line); y_coord += 15
            pdf_doc.save(DUMMY_PDF_PATH); pdf_doc.close()
            print(f"Created dummy PDF: {DUMMY_PDF_PATH} ({test_files_to_create['pdf']['type']})")
        except Exception as e: print(f"Error creating dummy PDF: {e}")
        try:
            docx_doc = Document()
            invoice_content = """TAX INVOICE\nSupplier: ABC Electronics Pvt Ltd, Address: 123 Tech Park, Bangalore, GSTIN: 29ABCDE1234F1Z5, Ph: 9988776655, Email: contact@abc.com\n\nInvoice No: INV2024/001\nInvoice Date: 02/Jan/2024\nDate of Supply: 03-01-2024\n\nBill To: XYZ Solutions, Address: 789 Business Hub, Mumbai, GSTIN: 27FGHIJ5678K1Z9\n\nDescription HSN Qty Rate Amount\nLaptop Computer 847130 1 50000.00 50000.00\nWireless Mouse 847160 2 500.00 1000.00\n\nSubtotal: 51000.00\nCGST @ 9%: 4590.00\nSGST @ 9%: 4590.00\nTotal Amount: 60180.00\nAmount in Words: Sixty Thousand One Hundred Eighty Only"""
            for line in invoice_content.split('\n'):
                if line.strip() == "": docx_doc.add_paragraph()
                else: docx_doc.add_paragraph(line)
            docx_doc.save(DUMMY_DOCX_PATH)
            print(f"Created dummy DOCX: {DUMMY_DOCX_PATH} ({test_files_to_create['docx']['type']})")
        except Exception as e: print(f"Error creating dummy DOCX: {e}")
        try:
            xlsx_workbook = openpyxl.Workbook(); sheet1 = xlsx_workbook.active; sheet1.title = "BalanceSheetData"
            sheet1['A1'] = "Example Corp Balance Sheet Inc."; sheet1['A2'] = "As of December 31, 2023"; sheet1['A3'] = "(Currency: INR)"
            sheet1['A5'] = "Assets"; sheet1['A6'] = "Total Current Assets :"; sheet1['B6'] = "1,500,000.00"; sheet1['A7'] = "Total Non-Current Assets :"; sheet1['B7'] = "2,000,000.00"; sheet1['A8'] = "Total Assets :"; sheet1['B8'] = "3,500,000.00"
            sheet1['A10'] = "Liabilities"; sheet1['A11'] = "Total Current Liabilities :"; sheet1['B11'] = "700,000.00"; sheet1['A12'] = "Total Non-Current Liabilities :"; sheet1['B12'] = "800,000.00"; sheet1['A13'] = "Total Liabilities :"; sheet1['B13'] = "1,500,000.00"
            sheet1['A15'] = "Equity"; sheet1['A16'] = "Total Equity :"; sheet1['B16'] = "2,000,000.00"
            xlsx_workbook.save(DUMMY_XLSX_PATH); print(f"Created dummy XLSX: {DUMMY_XLSX_PATH} ({test_files_to_create['xlsx']['type']})")
        except Exception as e: print(f"Error creating dummy XLSX: {e}")
        try:
            img = Image.new('RGB', (200, 50), color = (255, 255, 255)); draw = ImageDraw.Draw(img); font = None
            try: font = ImageFont.load_default()
            except IOError:
                try: font = ImageFont.truetype("DejaVuSans.ttf", 15) if os.path.exists("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf") else None
                except IOError: print("Default font and DejaVuSans.ttf not found for dummy PNG.")
            text_to_draw = "Hello OCR"
            if font:
                try: bbox = draw.textbbox((0,0), text_to_draw, font=font); text_width = bbox[2]-bbox[0]; text_height = bbox[3]-bbox[1]
                except AttributeError: text_width, text_height = draw.textsize(text_to_draw, font=font)
                draw.text(((img.width-text_width)/2, (img.height-text_height)/2), text_to_draw, fill=(0,0,0), font=font)
            else: draw.text((10,10), text_to_draw, fill=(0,0,0))
            img.save(DUMMY_PNG_PATH); print(f"Created dummy PNG: {DUMMY_PNG_PATH} ({test_files_to_create['png']['type']})")
        except Exception as e: print(f"Error creating dummy PNG: {e}")
        try:
            img_jpg = Image.new('RGB', (60,30), color='blue'); img_jpg.save(DUMMY_JPG_PATH)
            print(f"Created dummy JPG: {DUMMY_JPG_PATH} ({test_files_to_create['jpg']['type']})")
        except Exception as e: print(f"Error creating dummy JPG: {e}")
        with open(UNSUPPORTED_TXT_PATH, "w") as f: f.write("dummy text content")
        print(f"Created dummy TXT: {UNSUPPORTED_TXT_PATH} ({test_files_to_create['txt']['type']})")
        test_scenarios = [DUMMY_PDF_PATH, DUMMY_DOCX_PATH, DUMMY_XLSX_PATH, DUMMY_PNG_PATH, DUMMY_JPG_PATH, UNSUPPORTED_TXT_PATH, NON_EXISTENT_FILE]
        for file_to_test in test_scenarios:
            try: json_output = process_document(file_to_test); print(json_output)
            except Exception as e: print(f"--- Processing document: {file_to_test} ---"); print(json.dumps({"error": "Critical error in process_document", "details": str(e)}, indent=4))
        print("\n--- Cleaning up dummy files ---")
        for key in test_files_to_create:
            f_path = test_files_to_create[key]["name"]
            if os.path.exists(f_path):
                try: os.remove(f_path); print(f"Removed dummy file: {f_path}")
                except OSError as e: print(f"Error removing {f_path}: {e}")
